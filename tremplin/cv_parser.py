"""Extraction et analyse locale d'un CV (PDF, DOCX, texte)."""

from __future__ import annotations

import io
import re
from datetime import datetime

from .data_skills import SKILLS, SOFT_SKILLS
from .matching import (
    canonical_skill,
    detect_certifications,
    detect_education,
    detect_languages,
    extract_skills,
    normalize,
)

MAX_CHARS = 20000

ACTION_VERBS = [
    "analysé", "animé", "automatisé", "conçu", "construit", "coordonné", "créé",
    "développé", "dirigé", "encadré", "géré", "industrialisé", "lancé", "mis en place",
    "négocié", "optimisé", "organisé", "piloté", "préparé", "présenté", "réalisé",
    "rédigé", "structuré", "suivi", "testé", "transformé", "déployé", "mesuré",
]


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extrait le texte brut d'un CV, quel que soit le format supporté."""
    name = (filename or "").lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(file_bytes))
            pages = []
            for page in reader.pages[:10]:
                pages.append(page.extract_text() or "")
            text = "\n".join(pages)
        elif name.endswith(".docx"):
            from docx import Document

            document = Document(io.BytesIO(file_bytes))
            lines = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(lines)
        else:
            try:
                text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1", errors="ignore")
    except Exception as exc:  # pragma: no cover - dépend des fichiers utilisateur
        raise ValueError(f"Impossible de lire ce fichier : {exc}") from exc

    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_CHARS]


def _detect_identity(text: str) -> dict:
    identity = {"name": "", "email": "", "phone": "", "linkedin": "", "location": ""}
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if lines:
        candidate = lines[0]
        words = candidate.split()
        if 1 < len(words) <= 4 and not re.search(r"[\d@:/]", candidate) and len(candidate) <= 45:
            identity["name"] = candidate.title() if candidate.isupper() else candidate
    email = re.search(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}", text)
    if email:
        identity["email"] = email.group(0)
    phone = re.search(r"(?:(?:\+33|0)\s?[1-9])(?:[\s.\-]?\d{2}){4}", text)
    if phone:
        identity["phone"] = phone.group(0).strip()
    linkedin = re.search(r"(?:linkedin\.com/in/|linkedin\.com/pub/)[\w\-_%À-ÿ/]+", text, re.IGNORECASE)
    if linkedin:
        identity["linkedin"] = linkedin.group(0)
    cities = [
        ("casablanca", "Casablanca"), ("rabat", "Rabat"), ("marrakech", "Marrakech"),
        ("tanger", "Tanger"), ("fes", "Fès"), ("agadir", "Agadir"), ("oujda", "Oujda"),
        ("meknes", "Meknès"), ("kenitra", "Kénitra"), ("tetouan", "Tétouan"),
        ("el jadida", "El Jadida"), ("sale", "Salé"),
        ("paris", "Paris"), ("lyon", "Lyon"), ("marseille", "Marseille"),
        ("bordeaux", "Bordeaux"), ("lille", "Lille"), ("toulouse", "Toulouse"),
        ("nantes", "Nantes"), ("rennes", "Rennes"), ("strasbourg", "Strasbourg"),
        ("montpellier", "Montpellier"), ("nice", "Nice"), ("grenoble", "Grenoble"),
    ]
    normalized = normalize(text[:800])
    for city, display in cities:
        if re.search(r"(?<![a-z])" + city + r"(?![a-z])", normalized):
            identity["location"] = display
            break
    return identity


def _estimate_level(text: str, skill: str) -> int:
    """Estime un niveau prudent (3 par défaut) selon les indices textuels."""
    normalized = normalize(text)
    for alias in [skill] + list(SKILLS.get(skill, {}).get("aliases", []))[:4]:
        alias_norm = normalize(alias)
        if not alias_norm:
            continue
        for match in re.finditer(re.escape(alias_norm), normalized):
            window = normalized[max(0, match.start() - 60): match.end() + 60]
            if any(word in window for word in ("expert", "experte", "avance", "avancee", "5 ans", "4 ans", "3 ans", "maitrise")):
                return 4
            if any(word in window for word in ("notions", "debutant", "decouverte", "bases")):
                return 2
    return 3


def detect_profile_updates(text: str) -> dict:
    """Détecte les éléments du profil présents dans un CV."""
    if not text:
        return {}
    detected_skills = extract_skills(text)
    skills = {
        skill: _estimate_level(text, skill)
        for skill in detected_skills
        if skill in SKILLS and skill not in SOFT_SKILLS
    }
    qualities = [skill for skill in detected_skills if skill in SOFT_SKILLS]
    # Les compétences trop génériques citées une seule fois et sans niveau restent à 3.
    languages = detect_languages(text)
    certifications = detect_certifications(text)
    education = detect_education(text)
    identity = _detect_identity(text)

    sections = []
    normalized = normalize(text)
    for label, keywords in [
        ("Formation", ["formation", "diplome", "master", "licence"]),
        ("Expérience", ["experience", "stage", "emploi", "professionnel"]),
        ("Projets", ["projet", "projets", "realisation"]),
        ("Compétences", ["competences", "competences techniques", "skills"]),
        ("Langues", ["langues", "langue"]),
        ("Certifications", ["certification", "certifications"]),
    ]:
        if any(keyword in normalized for keyword in keywords):
            sections.append(label)

    return {
        "skills": skills,
        "qualities": qualities,
        "languages": languages,
        "certifications": certifications,
        "education": education,
        "identity": {key: value for key, value in identity.items() if value},
        "sections": sections,
    }


def audit_cv(text: str) -> dict:
    """Audit local du CV : structure, lisibilité ATS et impact."""
    if not text or len(text.strip()) < 80:
        return {
            "score": 0,
            "word_count": 0,
            "checks": [],
            "suggestions": ["Importe un CV lisible (PDF texte, DOCX) pour lancer l'audit."],
        }

    normalized = normalize(text)
    words = len(re.findall(r"\b[\w'’-]+\b", text))
    has_email = bool(re.search(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}", text))
    has_phone = bool(re.search(r"(?:(?:\+33|0)\s?[1-9])(?:[\s.\-]?\d{2}){4}", text))
    has_linkedin = "linkedin" in normalized
    has_formation = any(keyword in normalized for keyword in ("formation", "diplome", "master", "licence", "ecole"))
    has_experience = any(keyword in normalized for keyword in ("experience", "stage", "emploi"))
    has_skills = any(keyword in normalized for keyword in ("competences", "skills", "technologies"))
    has_languages = "langue" in normalized or "anglais" in normalized or "english" in normalized
    has_projects = "projet" in normalized
    quantified = len(re.findall(r"\d+\s?(?:%|€|k€|clients?|utilisateurs?|personnes|heures?|jours?|mois|ans?)", normalized))
    action_verbs = sum(1 for verb in ACTION_VERBS if verb in normalized)
    length_ok = 250 <= words <= 900

    checks = [
        {"label": "Longueur du CV (250 à 900 mots)", "passed": length_ok,
         "hint": "Un CV d'une page pour un stage : 350 à 600 mots suffisent."},
        {"label": "Coordonnées complètes (e-mail, téléphone)", "passed": has_email and has_phone,
         "hint": "Ajoute e-mail et téléphone cliquables en haut de page."},
        {"label": "Profil LinkedIn visible", "passed": has_linkedin,
         "hint": "Ajoute l'URL de ton profil LinkedIn, personnalisée si possible."},
        {"label": "Section Formation claire", "passed": has_formation,
         "hint": "Indique le diplôme, l'établissement et l'année d'obtention."},
        {"label": "Expériences ou stages détaillés", "passed": has_experience,
         "hint": "Même une expérience associative compte : décris tes missions."},
        {"label": "Section Compétences", "passed": has_skills,
         "hint": "Regroupe tes compétences par familles, les plus pertinentes d'abord."},
        {"label": "Langues renseignées", "passed": has_languages,
         "hint": "Précise le niveau CECR (B1, B2, C1) et un score si tu en as un."},
        {"label": "Projets valorisés", "passed": has_projects,
         "hint": "2 projets bien racontés valent mieux que 6 projets listés."},
        {"label": "Résultats quantifiés (au moins 2)", "passed": quantified >= 2,
         "hint": "Chiffre tes résultats : %, volumes, délais, budgets."},
        {"label": "Verbes d'action variés (au moins 4)", "passed": action_verbs >= 4,
         "hint": "Commence chaque puce par un verbe d'action au passé composé."},
    ]
    weights = {
        "Longueur du CV (250 à 900 mots)": 8,
        "Coordonnées complètes (e-mail, téléphone)": 10,
        "Profil LinkedIn visible": 6,
        "Section Formation claire": 10,
        "Expériences ou stages détaillés": 14,
        "Section Compétences": 12,
        "Langues renseignées": 8,
        "Projets valorisés": 12,
        "Résultats quantifiés (au moins 2)": 10,
        "Verbes d'action variés (au moins 4)": 10,
    }
    total = sum(weights.values())
    earned = sum(weights[c["label"]] for c in checks if c["passed"])
    score = int(round(100 * earned / total))

    suggestions: list[str] = []
    if not length_ok:
        suggestions.append("Vise 350 à 600 mots : supprime les répétitions et les mentions scolaires inutiles.")
    if quantified < 2:
        suggestions.append("Ajoute au moins deux résultats chiffrés pour crédibiliser tes missions.")
    if action_verbs < 4:
        suggestions.append("Remplace « participation à… » par des verbes forts : conçu, analysé, piloté…")
    if not has_linkedin:
        suggestions.append("Ajoute ton LinkedIn : c'est le premier réflexe de vérification d'un recruteur.")
    if not has_projects:
        suggestions.append("Crée une section Projets avec tes 2 réalisations les plus proches de l'offre visée.")
    if not has_languages:
        suggestions.append("Indique ton niveau de langue (B2, C1…) : c'est un filtre fréquent en stage.")
    if not suggestions:
        suggestions.append("CV solide ! Adapte maintenant chaque candidature aux mots-clés de l'offre.")

    return {"score": score, "word_count": words, "checks": checks, "suggestions": suggestions}
